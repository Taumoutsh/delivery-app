"""Delivery report export (PDF via reportlab, Word via python-docx)."""

import io
from datetime import datetime, timezone

from .db import build_variables, live_elapsed, render_command

STATUS_LABELS = {
    "todo": "À faire",
    "running": "En cours",
    "paused": "En pause",
    "done": "Validée",
    "skipped": "Non applicable",
}


class ExportUnavailable(Exception):
    pass


def fmt_duration(seconds):
    seconds = int(seconds or 0)
    h, rem = divmod(seconds, 3600)
    m = rem // 60
    if h:
        return f"{h} h {m:02d}"
    return f"{m} min"


def fmt_date(iso):
    if not iso:
        return "—"
    dt = datetime.strptime(iso, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=timezone.utc).astimezone()
    return dt.strftime("%d/%m/%Y %H:%M")


def _summary(delivery):
    tasks = delivery["tasks"]
    return {
        "variables": build_variables(delivery),
        "estimated": sum(t["estimated_minutes"] for t in tasks if t["status"] != "skipped") * 60,
        "spent": sum(live_elapsed(t) for t in tasks),
        "done": sum(1 for t in tasks if t["status"] == "done"),
        "skipped": sum(1 for t in tasks if t["status"] == "skipped"),
    }


def build_docx(delivery, db):
    try:
        from docx import Document
        from docx.shared import Cm, Pt, RGBColor
    except ImportError as exc:
        raise ExportUnavailable("python-docx n'est pas installé (pip install python-docx).") from exc

    s = _summary(delivery)
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    doc.add_heading(f"Rapport de livraison — {delivery['version']}", 0)
    if delivery["name"]:
        doc.add_paragraph(delivery["name"]).runs[0].italic = True

    info = doc.add_table(rows=0, cols=2)
    info.style = "Light List Accent 1"
    for label, value in [
        ("Version", delivery["version"]),
        ("Statut", "Clôturée" if delivery["status"] == "done" else "En cours"),
        ("Créée le", fmt_date(delivery["created_at"])),
        ("Clôturée le", fmt_date(delivery["completed_at"])),
        ("Progression", f"{delivery['progress']} % (pondérée)"),
        ("Tâches", f"{s['done']} validées, {s['skipped']} non applicables, {len(delivery['tasks'])} au total"),
        ("Temps", f"{fmt_duration(s['spent'])} passé / {fmt_duration(s['estimated'])} estimé"),
    ]:
        row = info.add_row().cells
        row[0].text, row[1].text = label, value

    if delivery["variables"]:
        doc.add_heading("Variables", 2)
        for key, value in delivery["variables"].items():
            doc.add_paragraph(f"{key} = {value}", style="List Bullet")

    doc.add_heading("Synthèse", 1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, ["#", "Tâche", "Statut", "Estimé", "Passé"]):
        cell.text = text
    for i, t in enumerate(delivery["tasks"], 1):
        cells = table.add_row().cells
        cells[0].text = str(i)
        cells[1].text = t["title"]
        cells[2].text = STATUS_LABELS.get(t["status"], t["status"])
        cells[3].text = fmt_duration(t["estimated_minutes"] * 60)
        cells[4].text = fmt_duration(live_elapsed(t))

    doc.add_heading("Détail des tâches", 1)
    for i, t in enumerate(delivery["tasks"], 1):
        doc.add_heading(f"{i}. {t['title']}", 2)
        meta = doc.add_paragraph()
        meta.add_run(
            f"{STATUS_LABELS.get(t['status'], t['status'])} · poids {t['weight']} · "
            f"estimé {fmt_duration(t['estimated_minutes'] * 60)} · passé {fmt_duration(live_elapsed(t))} · "
            f"démarrée {fmt_date(t['started_at'])} · terminée {fmt_date(t['completed_at'])}"
        ).font.color.rgb = RGBColor(0x55, 0x60, 0x6E)
        if t["description"]:
            doc.add_paragraph(render_command(t["description"], s["variables"]))
        if t["command"]:
            p = doc.add_paragraph()
            run = p.add_run(render_command(t["command"], s["variables"]))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        for c in t["comments"]:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"{fmt_date(c['created_at'])} — ").bold = True
            p.add_run(c["body"])
        for img in t["images"]:
            blob = db.get_image(img["id"])
            try:
                doc.add_picture(io.BytesIO(blob["data"]), width=Cm(14))
                doc.add_paragraph(img["filename"]).runs[0].italic = True
            except Exception:
                doc.add_paragraph(f"[image non intégrable : {img['filename']}]")

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def build_pdf(delivery, db):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.platypus import (Image, Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table,
                                        TableStyle)
    except ImportError as exc:
        raise ExportUnavailable("reportlab n'est pas installé (pip install reportlab).") from exc
    from xml.sax.saxutils import escape

    s = _summary(delivery)
    styles = getSampleStyleSheet()
    muted = ParagraphStyle("muted", parent=styles["Normal"], textColor=colors.HexColor("#55606E"), fontSize=8.5)
    code = ParagraphStyle("code", parent=styles["Code"], fontSize=8, backColor=colors.HexColor("#EEF2F5"),
                          borderPadding=6, leftIndent=6, rightIndent=6, spaceBefore=4, spaceAfter=8)
    accent = colors.HexColor("#1F5F7A")

    story = [Paragraph(f"Rapport de livraison — {escape(delivery['version'])}", styles["Title"])]
    if delivery["name"]:
        story.append(Paragraph(f"<i>{escape(delivery['name'])}</i>", styles["Normal"]))
    story.append(Spacer(1, 8))

    info = [
        ["Statut", "Clôturée" if delivery["status"] == "done" else "En cours"],
        ["Créée le", fmt_date(delivery["created_at"])],
        ["Clôturée le", fmt_date(delivery["completed_at"])],
        ["Progression", f"{delivery['progress']} % (pondérée)"],
        ["Tâches", f"{s['done']} validées, {s['skipped']} non applicables, {len(delivery['tasks'])} au total"],
        ["Temps", f"{fmt_duration(s['spent'])} passé / {fmt_duration(s['estimated'])} estimé"],
    ]
    info += [[f"{{{{{k}}}}}", v] for k, v in delivery["variables"].items()]
    t = Table(info, colWidths=[4 * cm, 12 * cm])
    t.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#D5DCE3")),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story += [t, Spacer(1, 14), Paragraph("Synthèse", styles["Heading2"])]

    rows = [["#", "Tâche", "Statut", "Estimé", "Passé"]]
    for i, task in enumerate(delivery["tasks"], 1):
        rows.append([str(i), Paragraph(escape(task["title"]), styles["Normal"]),
                     STATUS_LABELS.get(task["status"], task["status"]),
                     fmt_duration(task["estimated_minutes"] * 60), fmt_duration(live_elapsed(task))])
    t = Table(rows, colWidths=[1 * cm, 8.5 * cm, 3 * cm, 2 * cm, 2 * cm], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), accent),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6F8")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D5DCE3")),
    ]))
    story += [t, Spacer(1, 14), Paragraph("Détail des tâches", styles["Heading2"])]

    for i, task in enumerate(delivery["tasks"], 1):
        story.append(Paragraph(f"{i}. {escape(task['title'])}", styles["Heading3"]))
        story.append(Paragraph(escape(
            f"{STATUS_LABELS.get(task['status'], task['status'])} · poids {task['weight']} · "
            f"estimé {fmt_duration(task['estimated_minutes'] * 60)} · passé {fmt_duration(live_elapsed(task))} · "
            f"démarrée {fmt_date(task['started_at'])} · terminée {fmt_date(task['completed_at'])}"), muted))
        if task["description"]:
            text = escape(render_command(task["description"], s["variables"])).replace("\n", "<br/>")
            story.append(Paragraph(text, styles["Normal"]))
        if task["command"]:
            story.append(Preformatted(render_command(task["command"], s["variables"]), code))
        for c in task["comments"]:
            story.append(Paragraph(f"<b>{fmt_date(c['created_at'])}</b> — "
                                   f"{escape(c['body']).replace(chr(10), '<br/>')}", styles["Normal"]))
        for img in task["images"]:
            blob = db.get_image(img["id"])
            try:
                pic = Image(io.BytesIO(blob["data"]))
                ratio = min(1.0, (14 * cm) / pic.imageWidth, (10 * cm) / pic.imageHeight)
                pic.drawWidth, pic.drawHeight = pic.imageWidth * ratio, pic.imageHeight * ratio
                story += [Spacer(1, 4), pic, Paragraph(escape(img["filename"]), muted)]
            except Exception:
                story.append(Paragraph(escape(f"[image non intégrable : {img['filename']}]"), muted))
        story.append(Spacer(1, 8))

    out = io.BytesIO()
    SimpleDocTemplate(out, pagesize=A4, title=f"Livraison {delivery['version']}",
                      leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm, bottomMargin=1.8 * cm).build(story)
    return out.getvalue()
